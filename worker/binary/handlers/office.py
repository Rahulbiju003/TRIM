"""Office document handler — Tier 2 text extraction."""
from __future__ import annotations
import io
from pathlib import Path


def extract_text(file_path: str, raw_bytes: bytes) -> str | None:
    ext = Path(file_path).suffix.lower()
    try:
        if ext == ".docx":
            return _docx(raw_bytes)
        elif ext == ".xlsx":
            return _xlsx(raw_bytes)
        elif ext == ".pptx":
            return _pptx(raw_bytes)
        elif ext in (".odt", ".ods", ".odp"):
            return _odf(raw_bytes)
    except Exception:
        pass
    return None


def _docx(data: bytes) -> str | None:
    from docx import Document  # type: ignore[import-untyped]  # python-docx
    doc = Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs) or None


def _xlsx(data: bytes) -> str | None:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 200:
                rows.append(f"... ({ws.max_row - 200} more rows)")
                break
            vals = [str(c) if c is not None else "" for c in row]
            if any(v.strip() for v in vals):
                rows.append("\t".join(vals))
        if rows:
            parts.append(f"Sheet: {name}\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(parts) or None


def _pptx(data: bytes) -> str | None:
    from pptx import Presentation  # type: ignore[import-untyped]
    prs = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            slides.append(f"Slide {i + 1}:\n" + "\n".join(texts))
    return "\n\n".join(slides) or None


def _odf(data: bytes) -> str | None:
    try:
        from odf.opendocument import load as odf_load  # type: ignore[import-untyped]
        from odf.text import P  # type: ignore[import-untyped]
        doc = odf_load(io.BytesIO(data))
        texts = []
        for elem in doc.body.childNodes:
            for p in elem.getElementsByType(P):
                t = p.__str__().strip()
                if t:
                    texts.append(t)
        return "\n".join(texts) or None
    except ImportError:
        return None
